"""
Ground Truth Target Extraction: Binary Flood Indicator (`flooded`).

Derives the primary binary flood classification target `flooded` (0/1) based on
the HEC-RAS 2D passability target (`passability`):
- `0`: Passable (< 30 cm depth: passability 0 or 1)
- `1`: Flooded / Impassable (>= 30 cm depth: passability 2, representing standing water hazardous to vehicles)

Updates:
1. `feature_table.xlsx`: sheet `feature_table`, column `flooded`
2. `hai_chau_road_graph.graphml`: edge attribute `flooded`
3. `hai_chau_road_graph.gpkg`: column `flooded`
"""

from __future__ import annotations

import argparse
from pathlib import Path
import sys
import time

# Ensure UTF-8 output encoding on Windows PowerShell
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import geopandas as gpd
import networkx as nx
import numpy as np
import openpyxl
import pandas as pd

SCRIPT_DIR = Path(__file__).resolve().parent
WORKSPACE_DIR = SCRIPT_DIR.parent.parent

DEFAULT_EXCEL_PATH = WORKSPACE_DIR / "dataset" / "feature_table.xlsx"
DEFAULT_GRAPHML_PATH = WORKSPACE_DIR / "dataset" / "hai_chau_road_graph.graphml"
DEFAULT_GPKG_PATH = WORKSPACE_DIR / "dataset" / "hai_chau_road_graph.gpkg"
DEFAULT_OUT_CSV = SCRIPT_DIR / "flooded_target_summary.csv"


def derive_flooded_target(
    excel_path: Path,
    impassable_only: bool = True,
) -> tuple[list[str], list[int], pd.Series]:
    """
    Load passability from feature_table.xlsx and derive binary flooded target.
    """
    if not excel_path.exists():
        raise FileNotFoundError(f"feature_table.xlsx not found at: {excel_path}")

    print(f"\n[Step 1/4] Loading passability labels from {excel_path.name}...")
    df = pd.read_excel(excel_path, sheet_name="feature_table")
    if "passability" not in df.columns:
        raise KeyError("'passability' column missing from feature_table sheet.")

    segment_ids = df["segment_id"].astype(str).tolist()
    passability = df["passability"].fillna(0).astype(int)

    print("\n[Step 2/4] Deriving binary flood target labels...")
    if impassable_only:
        # Primary benchmark definition: impassable >= 30 cm (passability == 2)
        flooded = (passability == 2).astype(int).tolist()
        criterion_desc = "Impassable threshold (passability == 2, depth >= 30 cm)"
    else:
        # Broad flood definition: any flood >= 10 cm (passability >= 1)
        flooded = (passability >= 1).astype(int).tolist()
        criterion_desc = "Broad flood threshold (passability >= 1, depth >= 10 cm)"

    print(f"Criterion: {criterion_desc}")
    return segment_ids, flooded, passability


def update_excel_flooded(excel_path: Path, segment_ids: list[str], flooded: list[int]) -> None:
    """
    Update 'flooded' column in 'feature_table' sheet of feature_table.xlsx.
    """
    print(f"\n[Step 3/4] Updating Excel spreadsheet: {excel_path.name}...")
    wb = openpyxl.load_workbook(excel_path)
    if "feature_table" not in wb.sheetnames:
        raise ValueError(f"'feature_table' sheet missing from {excel_path}")

    ws = wb["feature_table"]
    header_row = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    col_map = {name: idx + 1 for idx, name in enumerate(header_row) if name is not None}

    if "flooded" not in col_map:
        raise KeyError(f"'flooded' column missing from headers: {header_row}")
    if "segment_id" not in col_map:
        raise KeyError(f"'segment_id' column missing from headers: {header_row}")

    flood_col_idx = col_map["flooded"]
    seg_col_idx = col_map["segment_id"]
    flood_lookup = dict(zip(segment_ids, flooded))

    updated = 0
    for r in range(2, ws.max_row + 1):
        sid = str(ws.cell(r, seg_col_idx).value)
        if sid in flood_lookup:
            ws.cell(r, flood_col_idx).value = int(flood_lookup[sid])
            updated += 1

    wb.save(excel_path)
    print(f"Successfully updated 'flooded' column for {updated:,} rows in {excel_path.name}.")


def update_graphml_flooded(graphml_path: Path, segment_ids: list[str], flooded: list[int]) -> None:
    """Encode flooded attribute into GraphML edges."""
    if not graphml_path.exists():
        print(f"GraphML not found at {graphml_path}, skipping GraphML update.")
        return

    print(f"\n[Step 4a/4] Updating GraphML edge attributes: {graphml_path.name}...")
    G = nx.read_graphml(graphml_path)
    flood_lookup = dict(zip(segment_ids, flooded))

    updated = 0
    for u, v, k, d in G.edges(keys=True, data=True):
        sid = d.get("segment_id")
        if sid and sid in flood_lookup:
            d["flooded"] = int(flood_lookup[sid])
            updated += 1

    nx.write_graphml(G, graphml_path)
    print(f"Encoded 'flooded' attribute into {updated:,} edges in {graphml_path.name}.")


def update_gpkg_flooded(gpkg_path: Path, segment_ids: list[str], flooded: list[int]) -> None:
    """Update flooded column in road GeoPackage layer."""
    if not gpkg_path.exists():
        print(f"GPKG not found at {gpkg_path}, skipping GPKG update.")
        return

    print(f"\n[Step 4b/4] Updating GeoPackage: {gpkg_path.name}...")
    gdf = gpd.read_file(gpkg_path)
    feat_df = pd.DataFrame({
        "segment_id": segment_ids,
        "flooded": flooded
    })

    if "flooded" in gdf.columns:
        gdf = gdf.drop(columns=["flooded"])

    merged = gdf.merge(feat_df, on="segment_id", how="left")
    merged.to_file(gpkg_path, driver="GPKG")
    print(f"Saved updated GeoPackage with 'flooded' column to {gpkg_path.name}.")


def main():
    parser = argparse.ArgumentParser(description="Derive binary flooded target and update feature_table.xlsx.")
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL_PATH, help="Path to feature_table.xlsx")
    parser.add_argument("--graphml", type=Path, default=DEFAULT_GRAPHML_PATH, help="Path to hai_chau_road_graph.graphml")
    parser.add_argument("--gpkg", type=Path, default=DEFAULT_GPKG_PATH, help="Path to hai_chau_road_graph.gpkg")
    parser.add_argument("--broad-flood", action="store_true", help="Use broad flood threshold (depth >= 10 cm)")
    parser.add_argument("--out-csv", type=Path, default=DEFAULT_OUT_CSV, help="Path to output summary CSV")
    args = parser.parse_args()

    t_start = time.time()
    print("=" * 80)
    print("GROUND TRUTH TARGET EXTRACTION: Binary Flooded Indicator (flooded)")
    print("=" * 80)
    print(f"Excel File  : {args.excel}")
    print(f"GraphML File: {args.graphml}")
    print(f"GPKG File   : {args.gpkg}")
    print(f"Broad Flood : {args.broad_flood}")
    print("-" * 80)

    # 1 & 2. Derive flooded
    segment_ids, flooded, passability = derive_flooded_target(args.excel, impassable_only=not args.broad_flood)

    # 3. Update Excel
    update_excel_flooded(args.excel, segment_ids, flooded)

    # 4. Update GraphML & GPKG
    update_graphml_flooded(args.graphml, segment_ids, flooded)
    update_gpkg_flooded(args.gpkg, segment_ids, flooded)

    # 5. Output summary
    arr = np.array(flooded)
    total = len(arr)
    pos = int((arr == 1).sum())
    neg = int((arr == 0).sum())
    print("\n" + "=" * 80)
    print("BINARY FLOOD TARGET SUMMARY:")
    print("=" * 80)
    print(f"  [0] Not Flooded (Passable)   : {neg:5,} ({neg / total * 100:5.1f}%)")
    print(f"  [1] Flooded (Impassable)     : {pos:5,} ({pos / total * 100:5.1f}%)")
    print(f"  Total Road Segments          : {total:5,} (100.0%)")
    print("=" * 80)

    summary_df = pd.DataFrame({
        "segment_id": segment_ids,
        "passability": passability,
        "flooded": flooded
    })
    args.out_csv.parent.mkdir(parents=True, exist_ok=True)
    summary_df.to_csv(args.out_csv, index=False)
    print(f"Saved summary CSV to: {args.out_csv}")

    elapsed = time.time() - t_start
    print(f"\nAll operations completed successfully in {elapsed:.2f} seconds.")


if __name__ == "__main__":
    main()
