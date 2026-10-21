"""
Network Feature Extraction: Edge Betweenness Centrality (`betweenness`).

Computes normalized edge betweenness centrality on the Hai Chau road graph
weighted by segment length (`length_m`) using Dijkstra shortest paths.

Updates:
1. `feature_table.xlsx`: sheet `feature_table`, column `betweenness`
2. `hai_chau_road_graph.graphml`: edge attribute `betweenness`
3. `hai_chau_road_graph.gpkg`: column `betweenness`
"""

from __future__ import annotations

import argparse
from pathlib import Path
import sys
import time

# Ensure UTF-8 output encoding on Windows PowerShell
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import geopandas as gpd
import networkx as nx
import numpy as np
import openpyxl
import pandas as pd

SCRIPT_DIR = Path(__file__).resolve().parent
WORKSPACE_DIR = SCRIPT_DIR.parent.parent

DEFAULT_EXCEL_PATH = WORKSPACE_DIR / "dataset" / "feature_table.xlsx"
DEFAULT_GRAPHML_PATH = WORKSPACE_DIR / "dataset" / "hai_chau_road_graph.graphml"
DEFAULT_GPKG_PATH = WORKSPACE_DIR / "dataset" / "hai_chau_road_graph.gpkg"
DEFAULT_OUT_CSV = SCRIPT_DIR / "betweenness_summary.csv"


def compute_edge_betweenness(graphml_path: Path) -> tuple[nx.MultiDiGraph, dict[str, float]]:
    """
    Load GraphML road network, assign edge weights based on length_m,
    and compute normalized edge betweenness centrality.
    """
    if not graphml_path.exists():
        raise FileNotFoundError(f"Road graph GraphML not found at: {graphml_path}")

    print(f"\n[Step 1/4] Loading road graph from {graphml_path.name}...")
    G = nx.read_graphml(graphml_path)
    print(f"Loaded road graph: {G.number_of_nodes():,} nodes, {G.number_of_edges():,} edges.")

    # Assign length weight
    for u, v, k, d in G.edges(keys=True, data=True):
        d["weight"] = float(d.get("length_m", 1.0))

    print("\n[Step 2/4] Computing normalized edge betweenness centrality (Dijkstra routing)...")
    t0 = time.time()
    ebc = nx.edge_betweenness_centrality(G, weight="weight", normalized=True)
    elapsed = time.time() - t0
    print(f"Betweenness computation completed in {elapsed:.2f} seconds.")

    # Build lookup by segment_id
    ebc_lookup = {}
    for (u, v, k), val in ebc.items():
        sid = G[u][v][k].get("segment_id")
        if sid:
            ebc_lookup[sid] = float(val)

    return G, ebc_lookup


def update_excel_betweenness(excel_path: Path, ebc_lookup: dict[str, float]) -> list[str]:
    """
    Update 'betweenness' column in 'feature_table' sheet of feature_table.xlsx.
    """
    print(f"\n[Step 3/4] Updating Excel spreadsheet: {excel_path.name}...")
    wb = openpyxl.load_workbook(excel_path)
    if "feature_table" not in wb.sheetnames:
        raise ValueError(f"'feature_table' sheet missing from {excel_path}")

    ws = wb["feature_table"]
    header_row = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    col_map = {name: idx + 1 for idx, name in enumerate(header_row) if name is not None}

    if "betweenness" not in col_map:
        raise KeyError(f"'betweenness' column missing from headers: {header_row}")
    if "segment_id" not in col_map:
        raise KeyError(f"'segment_id' column missing from headers: {header_row}")

    bet_col_idx = col_map["betweenness"]
    seg_col_idx = col_map["segment_id"]

    updated = 0
    segment_ids = []
    for r in range(2, ws.max_row + 1):
        sid = ws.cell(r, seg_col_idx).value
        if sid is not None:
            segment_ids.append(str(sid))
            val = ebc_lookup.get(str(sid), 0.0)
            ws.cell(r, bet_col_idx).value = round(float(val), 8)
            updated += 1

    wb.save(excel_path)
    print(f"Successfully updated 'betweenness' column for {updated:,} rows in {excel_path.name}.")
    return segment_ids


def update_graphml_betweenness(G: nx.MultiDiGraph, graphml_path: Path, ebc_lookup: dict[str, float]) -> None:
    """Encode betweenness attribute into GraphML edges."""
    print(f"\n[Step 4a/4] Updating GraphML edge attributes: {graphml_path.name}...")
    updated = 0
    for u, v, k, d in G.edges(keys=True, data=True):
        sid = d.get("segment_id")
        if sid and sid in ebc_lookup:
            d["betweenness"] = float(ebc_lookup[sid])
            updated += 1
    nx.write_graphml(G, graphml_path)
    print(f"Encoded 'betweenness' attribute into {updated:,} edges in {graphml_path.name}.")


def update_gpkg_betweenness(gpkg_path: Path, segment_ids: list[str], ebc_lookup: dict[str, float]) -> None:
    """Update betweenness column in road GeoPackage layer."""
    if not gpkg_path.exists():
        print(f"GPKG not found at {gpkg_path}, skipping GPKG update.")
        return

    print(f"\n[Step 4b/4] Updating GeoPackage: {gpkg_path.name}...")
    gdf = gpd.read_file(gpkg_path)
    feat_df = pd.DataFrame({
        "segment_id": segment_ids,
        "betweenness": [ebc_lookup.get(sid, 0.0) for sid in segment_ids]
    })

    if "betweenness" in gdf.columns:
        gdf = gdf.drop(columns=["betweenness"])

    merged = gdf.merge(feat_df, on="segment_id", how="left")
    merged.to_file(gpkg_path, driver="GPKG")
    print(f"Saved updated GeoPackage with 'betweenness' column to {gpkg_path.name}.")


def main():
    parser = argparse.ArgumentParser(description="Compute edge betweenness centrality and update feature_table.xlsx.")
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL_PATH, help="Path to feature_table.xlsx")
    parser.add_argument("--graphml", type=Path, default=DEFAULT_GRAPHML_PATH, help="Path to hai_chau_road_graph.graphml")
    parser.add_argument("--gpkg", type=Path, default=DEFAULT_GPKG_PATH, help="Path to hai_chau_road_graph.gpkg")
    parser.add_argument("--out-csv", type=Path, default=DEFAULT_OUT_CSV, help="Path to output summary CSV")
    args = parser.parse_args()

    t_start = time.time()
    print("=" * 80)
    print("ROAD NETWORK FEATURE EXTRACTION: Edge Betweenness Centrality")
    print("=" * 80)
    print(f"GraphML File: {args.graphml}")
    print(f"Excel File  : {args.excel}")
    print(f"GPKG File   : {args.gpkg}")
    print("-" * 80)

    # 1 & 2. Compute betweenness
    G, ebc_lookup = compute_edge_betweenness(args.graphml)

    # 3. Update Excel
    segment_ids = update_excel_betweenness(args.excel, ebc_lookup)

    # 4. Update GraphML & GPKG
    update_graphml_betweenness(G, args.graphml, ebc_lookup)
    update_gpkg_betweenness(args.gpkg, segment_ids, ebc_lookup)

    # 5. Output summary CSV and statistics
    values = [ebc_lookup.get(sid, 0.0) for sid in segment_ids]
    arr = np.array(values)
    print("\n" + "=" * 80)
    print("BETWEENNESS CENTRALITY SUMMARY:")
    print("=" * 80)
    print(f"  Count    : {len(arr):,}")
    print(f"  Min      : {arr.min():.8f}")
    print(f"  Max      : {arr.max():.8f}")
    print(f"  Mean     : {arr.mean():.8f}")
    print(f"  Median   : {np.median(arr):.8f}")
    print(f"  Non-zero : {(arr > 0).sum():,} ({(arr > 0).sum() / len(arr) * 100:.1f}%)")
    print("=" * 80)

    summary_df = pd.DataFrame({
        "segment_id": segment_ids,
        "betweenness": values
    }).sort_values("betweenness", ascending=False)

    args.out_csv.parent.mkdir(parents=True, exist_ok=True)
    summary_df.to_csv(args.out_csv, index=False)
    print(f"Saved summary CSV to: {args.out_csv}")

    elapsed = time.time() - t_start
    print(f"\nAll operations completed successfully in {elapsed:.2f} seconds.")


if __name__ == "__main__":
    main()
